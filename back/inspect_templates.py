import openpyxl
from pptx import Presentation
import docx
import os

templates_dir = "/app/app/assets/templates"

excel_path = os.path.join(templates_dir, "plantilla excel.xlsx")
word_path = os.path.join(templates_dir, "plantilla-word.docx")
pptx_path = os.path.join(templates_dir, "plantilla-powerpoint.pptx")

print("--- EXCEL ---")
if os.path.exists(excel_path):
    wb = openpyxl.load_workbook(excel_path)
    print("Sheets:", wb.sheetnames)
    ws = wb.active
    print("Active sheet cells:")
    for r in range(1, 15):
        row_vals = [ws.cell(r, c).value for c in range(1, 10)]
        if any(row_vals):
            print(f"Row {r}:", row_vals)
else:
    print("Excel template not found:", excel_path)

print("\n--- WORD ---")
if os.path.exists(word_path):
    doc = docx.Document(word_path)
    print("Paragraphs count:", len(doc.paragraphs))
    for i, p in enumerate(doc.paragraphs[:10]):
        if p.text.strip():
            print(f"Paragraph {i}:", p.text)
else:
    print("Word template not found:", word_path)

print("\n--- PPTX ---")
if os.path.exists(pptx_path):
    prs = Presentation(pptx_path)
    print("Slides count:", len(prs.slides))
    print("Slide layouts:")
    for i, l in enumerate(prs.slide_layouts):
        print(f"Layout {i}: {l.name}")
else:
    print("PPTX template not found:", pptx_path)
