from pptx import Presentation
import os

pptx_path = "/app/app/assets/templates/plantilla-powerpoint.pptx"

if os.path.exists(pptx_path):
    prs = Presentation(pptx_path)
    print("Slides count:", len(prs.slides))
    for idx, slide in enumerate(prs.slides):
        if idx >= 15:
            break
        print(f"\n--- Slide {idx} ---")
        print("Layout:", slide.slide_layout.name)
        # Inspect shapes on the slide
        for s_idx, shape in enumerate(slide.shapes):
            text = getattr(shape, "text", "")
            shape_type = shape.shape_type
            print(f"  Shape {s_idx}: {shape.name} (type: {shape_type})")
            if text.strip():
                print(f"    Text: {text.strip()}")
else:
    print("PPTX template not found")
