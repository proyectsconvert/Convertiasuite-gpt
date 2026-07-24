import os
import json
import requests
from concurrent.futures import ThreadPoolExecutor, as_completed

SUPABASE_URL = os.environ["SUPABASE_URL"].rstrip("/")
SUPABASE_SERVICE_KEY = os.environ["SUPABASE_SERVICE_KEY"]

OLLAMA_HOST = os.environ.get("OLLAMA_BASE_URL", "http://localhost:11434")
OLLAMA_EMBED_MODEL = os.environ.get("OLLAMA_EMBED_MODEL", "nomic-embed-text")

CF_ACCESS_CLIENT_ID = os.environ.get("CF_ACCESS_CLIENT_ID")
CF_ACCESS_CLIENT_SECRET = os.environ.get("CF_ACCESS_CLIENT_SECRET")

BUCKET_NAME = "ai_files"
CHARS_PER_TOKEN = 4
MAX_TOKENS = 500
OVERLAP = 50


def supabase_headers(extra=None):
    headers = {
        "apikey": SUPABASE_SERVICE_KEY,
        "Authorization": f"Bearer {SUPABASE_SERVICE_KEY}",
        "Content-Type": "application/json",
    }
    if extra:
        headers.update(extra)
    return headers


def ollama_headers():
    headers = {}
    if CF_ACCESS_CLIENT_ID and CF_ACCESS_CLIENT_SECRET:
        headers["CF-Access-Client-Id"] = CF_ACCESS_CLIENT_ID
        headers["CF-Access-Client-Secret"] = CF_ACCESS_CLIENT_SECRET
    return headers


# Supabase: leer pendientes / escribir chunks (vía REST, no SDK)
def get_pending_sources():
    url = f"{SUPABASE_URL}/rest/v1/rpc/get_pending_rag_sources"
    response = requests.post(url, headers=supabase_headers(), json={}, timeout=30)
    response.raise_for_status()
    return response.json()


def delete_existing_chunks(source_id):
    url = f"{SUPABASE_URL}/rest/v1/rag_chunks"
    response = requests.delete(
        url,
        headers=supabase_headers(),
        params={"source_id": f"eq.{source_id}"},
        timeout=30,
    )
    response.raise_for_status()


def insert_chunks(rows):
    if not rows:
        return
    url = f"{SUPABASE_URL}/rest/v1/rag_chunks"
    response = requests.post(
        url,
        headers=supabase_headers({"Prefer": "return=minimal"}),
        json=rows,
        timeout=60,
    )
    response.raise_for_status()


def download_from_storage(storage_path):
    clean_path = storage_path
    bucket = BUCKET_NAME

    if "/" in storage_path:
        first_part = storage_path.split("/")[0]
        if first_part in ("ai_files", "attachments"):
            bucket = first_part
            clean_path = storage_path[len(first_part) + 1:]

    url = f"{SUPABASE_URL}/storage/v1/object/{bucket}/{clean_path}"
    response = requests.get(url, headers=supabase_headers(), timeout=60)
    response.raise_for_status()
    return response.content


# Ollama: embeddings
def embed_batch(texts):
    url = f"{OLLAMA_HOST}/api/embed"
    response = requests.post(
        url,
        json={"model": OLLAMA_EMBED_MODEL, "input": texts},
        headers=ollama_headers(),
        timeout=180,
    )
    response.raise_for_status()
    return response.json()["embeddings"]


def split_into_chunks(text, max_tokens=MAX_TOKENS, overlap=OVERLAP):
    max_chars = max_tokens * CHARS_PER_TOKEN
    overlap_chars = overlap * CHARS_PER_TOKEN

    text = (text or "").strip()
    if not text:
        return []
    if len(text) <= max_chars:
        return [text]

    chunks = []
    start = 0
    while start < len(text):
        end = start + max_chars
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        start = end - overlap_chars
    return chunks


def table_to_text(table, index):
    headers = table.get("headers") or []
    rows = table.get("rows") or []
    if not headers and not rows:
        return ""

    lines = []
    if headers:
        lines.append(" | ".join(headers))
        lines.append(" | ".join(["---"] * len(headers)))
    for row in rows:
        lines.append(" | ".join(str(c) for c in row))

    name = table.get("name") or f"tabla_{index}"
    return f"Tabla: {name}\n" + "\n".join(lines)


def chunks_from_parsed_content(parsed_content):
    if isinstance(parsed_content, str):
        parsed_content = json.loads(parsed_content)

    results = []

    sections = parsed_content.get("sections") or []
    if sections:
        for section in sections:
            content = (section.get("content") or "").strip()
            if not content:
                continue
            base_meta = {
                "type": "text",
                "page": (section.get("metadata") or {}).get("page"),
                "section_title": section.get("title"),
            }
            for i, sub in enumerate(split_into_chunks(content)):
                meta = dict(base_meta)
                if i > 0:
                    meta["sub_chunk"] = i
                results.append((sub, meta))
    else:
        for sub in split_into_chunks(parsed_content.get("text", "")):
            results.append((sub, {"type": "text"}))

    for idx, table in enumerate(parsed_content.get("tables") or []):
        table_text = table_to_text(table, idx)
        if table_text:
            results.append((table_text, {"type": "table", "table_index": idx, "table_name": table.get("name")}))

    return results


def extract_text_from_file(file_bytes, filename):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext == "pdf":
        from pypdf import PdfReader
        import io
        reader = PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == "docx":
        import docx
        import io
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    if ext in ("xlsx", "xls"):
        import openpyxl
        import io
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                lines.append(" | ".join(str(c) for c in row if c is not None))
        return "\n".join(lines)

    return file_bytes.decode("utf-8", errors="ignore")


# Orquestación por documento
def get_chunks_for_row(row):
    if row.get("parsed_content"):
        return chunks_from_parsed_content(row["parsed_content"])

    if row.get("extracted_text"):
        return [(c, {"type": "text"}) for c in split_into_chunks(row["extracted_text"])]

    if row.get("storage_path"):
        file_bytes = download_from_storage(row["storage_path"])
        text = extract_text_from_file(file_bytes, row["file_name"])
        return [(c, {"type": "text"}) for c in split_into_chunks(text)]

    raise ValueError(f"Sin contenido disponible para {row['source_type']}:{row['external_ref']}")


def ingest_one(row):
    source_id = row["source_id"]
    chunk_pairs = get_chunks_for_row(row)

    if not chunk_pairs:
        print(f"[SKIP] {row['source_type']}:{row['file_name']} -> sin contenido extraíble")
        return

    texts = [t for t, _ in chunk_pairs]
    vectors = embed_batch(texts)

    delete_existing_chunks(source_id)

    rows = [
        {
            "source_id": source_id,
            "chunk_index": i,
            "content": text,
            "metadata": meta,
            "embedding": vector,
            "source_updated_at": row["source_updated_at"],
        }
        for i, ((text, meta), vector) in enumerate(zip(chunk_pairs, vectors))
    ]
    insert_chunks(rows)

    print(f"[OK] {row['source_type']}:{row['file_name']} -> {len(rows)} chunks")


INGEST_WORKERS = int(os.environ.get("RAG_INGEST_WORKERS", 4))


def main():
    print("Buscando documentos pendientes de ingesta...")
    pending = get_pending_sources()
    print(f"{len(pending)} documentos pendientes de (re)procesar")

    errors = 0
    with ThreadPoolExecutor(max_workers=INGEST_WORKERS) as executor:
        futures = {executor.submit(ingest_one, row): row for row in pending}

        for future in as_completed(futures):
            row = futures[future]
            try:
                future.result()
            except Exception as exc:
                errors += 1
                print(f"[ERROR] {row['source_type']}:{row.get('file_name')} -> {exc}")

    if errors:
        print(f"Terminado con {errors} errores")
        exit(1)

    print("Ingesta completada sin errores.")


if __name__ == "__main__":
    main()